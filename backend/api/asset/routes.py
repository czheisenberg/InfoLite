# -*- coding: utf-8 -*-
# @Author: Info Lite
# @Desc: 资产相关API路由
from fastapi import APIRouter, Query, Depends, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse
from io import BytesIO
import time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from scan_nmap import check_ip_allowed as check_ip_allowed_nmap, ip_full_scan as ip_full_scan_nmap, custom_nmap_scan
from scan_core import check_ip_allowed as check_ip_allowed_socket, ip_full_scan as ip_full_scan_socket
from db_operate import query_asset_from_db, insert_asset_to_mysql, insert_asset_to_es
from api.deps import get_db

router = APIRouter(prefix="/api/asset", tags=["资产操作"])


@router.get("/query", summary="资产查询/扫描接口")
async def query_asset(
    ip: str = Query(..., description="目标IP地址，如127.0.0.1"),
    refresh: bool = Query(False, description="是否强制刷新扫描，默认False"),
    port_option: str = Query("common", description="端口选项：common（常用端口）或custom（自定义端口范围）"),
    port_range: str = Query("1-65535", description="自定义端口范围，如1-65535"),
    scan_mode: str = Query("nmap", description="扫描引擎：nmap 或 socket"),
    current_user: dict = Depends(lambda: None),
    db: dict = Depends(get_db)
):
    """
    资产查询/扫描核心接口：
    1. 非刷新+IP有有效数据 → 直接返回数据库数据
    2. 非刷新+无数据/数据过期 → 触发扫描+入库+返回
    3. 刷新 → 强制触发扫描+覆盖数据库+返回
    """
    try:
        # 校验扫描模式
        if scan_mode not in ["nmap", "socket"]:
            scan_mode = "nmap"

        # 选择对应扫描引擎
        if scan_mode == "nmap":
            check_ip_allowed_func = check_ip_allowed_nmap
            ip_full_scan_func = ip_full_scan_nmap
        else:
            check_ip_allowed_func = check_ip_allowed_socket
            ip_full_scan_func = ip_full_scan_socket

        # 第一步：IP格式简单校验
        ip_parts = ip.split(".")
        if len(ip_parts) != 4 or not all(part.isdigit() and 0 <= int(part) <= 255 for part in ip_parts):
            raise HTTPException(status_code=400, detail=f"IP地址格式错误：{ip}")
        
        # 第二步：IP白名单校验
        if not check_ip_allowed_func(ip, db["scan_config"]["allowed_ip"]):
            raise HTTPException(status_code=403, detail=f"IP {ip} 不在扫描白名单内，禁止扫描！")
        
        # 第三步：非刷新模式，先查询数据库
        if not refresh:
            db_result = query_asset_from_db(ip, db["mysql_conn"], db["mysql_cursor"], db["es_conn"], db["scan_expire_time"])
            if db_result["is_valid"] and db_result["data"]:
                return JSONResponse(
                    status_code=200,
                    content={
                        "code": 200,
                        "msg": "查询成功（从数据库获取）",
                        "data": db_result["data"],
                        "scan_type": "db_query",
                        "scan_mode": scan_mode
                    }
                )
        
        # 第四步：触发扫描
        scan_config = db["scan_config"].copy()
        scan_config["port_option"] = port_option
        scan_config["port_range"] = port_range
        asset_list = ip_full_scan_func(ip, scan_config)
        if not asset_list:
            return JSONResponse(
                status_code=200,
                content={
                    "code": 200,
                    "msg": f"IP {ip} 无开放端口或扫描失败",
                    "data": [],
                    "scan_type": "scan",
                    "scan_mode": scan_mode
                }
            )
        
        # 第五步：扫描结果入库
        insert_asset_to_mysql(db["mysql_conn"], db["mysql_cursor"], asset_list)
        insert_asset_to_es(db["es_conn"], asset_list)
        
        # 第六步：返回扫描结果
        return JSONResponse(
            status_code=200,
            content={
                "code": 200,
                "msg": "扫描成功（新结果已入库）",
                "data": asset_list,
                "scan_type": "scan",
                "scan_mode": scan_mode
            }
        )
    
    except HTTPException as e:
        return JSONResponse(
            status_code=e.status_code,
            content={"code": e.status_code, "msg": e.detail, "data": []}
        )
    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={"code": 500, "msg": f"服务器内部错误：{str(e)}", "data": []}
        )


@router.get("/nmap-scan", summary="Nmap自定义扫描接口")
async def nmap_custom_scan(
    ip: str = Query(..., description="目标IP地址"),
    ports: str = Query("", description="端口范围，如 1-65535 或 80,443"),
    scan_args: str = Query("", description="Nmap扫描参数，如 -sV -T4 -O"),
    save_to_db: bool = Query(True, description="是否将结果保存到数据库"),
    current_user: dict = Depends(lambda: None),
    db: dict = Depends(get_db)
):
    """
    Nmap自定义扫描接口 - 用户手动指定扫描参数
    """
    try:
        ip_parts = ip.split(".")
        if len(ip_parts) != 4 or not all(part.isdigit() and 0 <= int(part) <= 255 for part in ip_parts):
            raise HTTPException(status_code=400, detail=f"IP地址格式错误：{ip}")

        if not check_ip_allowed_nmap(ip, db["scan_config"]["allowed_ip"]):
            raise HTTPException(status_code=403, detail=f"IP {ip} 不在扫描白名单内，禁止扫描！")

        if not scan_args.strip():
            scan_args = "-sV -T4"

        asset_list = custom_nmap_scan(
            ip=ip,
            scan_config=db["scan_config"],
            custom_args=scan_args,
            ports=ports
        )

        if save_to_db and asset_list:
            insert_asset_to_mysql(db["mysql_conn"], db["mysql_cursor"], asset_list)
            insert_asset_to_es(db["es_conn"], asset_list)

        return JSONResponse(
            status_code=200,
            content={
                "code": 200,
                "msg": "Nmap自定义扫描成功",
                "data": asset_list,
                "scan_type": "custom_nmap",
                "scan_args": scan_args,
                "target_ip": ip,
                "ports": ports
            }
        )

    except HTTPException as e:
        return JSONResponse(
            status_code=e.status_code,
            content={"code": e.status_code, "msg": e.detail, "data": []}
        )
    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={"code": 500, "msg": f"服务器内部错误：{str(e)}", "data": []}
        )


@router.get("/export-excel", summary="导出资产Excel接口")
async def export_asset_excel(
    ip: str = Query(..., description="目标IP地址"),
    current_user: dict = Depends(lambda: None),
    db: dict = Depends(get_db)
):
    """
    导出指定IP的扫描结果为Excel文件
    """
    try:
        # 查询数据库中的资产数据
        db_result = query_asset_from_db(ip, db["mysql_conn"], db["mysql_cursor"], db["es_conn"], db["scan_expire_time"])
        asset_list = db_result["data"] if db_result["is_valid"] else []

        if not asset_list:
            raise HTTPException(status_code=404, detail=f"IP {ip} 没有可导出的扫描数据")

        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "资产扫描结果"

        # 设置表头样式
        header_font = Font(name='微软雅黑', bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # 边框样式
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 写入表头
        headers = ["IP地址", "端口", "协议", "状态", "服务", "版本", "产品", "其他信息", "扫描时间"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写入数据
        data_font = Font(name='微软雅黑', size=10)
        data_alignment = Alignment(horizontal="center", vertical="center")

        for row, asset in enumerate(asset_list, 2):
            row_data = [
                asset.get("ip", ""),
                asset.get("port", ""),
                asset.get("protocol", ""),
                asset.get("state", ""),
                asset.get("service", ""),
                asset.get("version", ""),
                asset.get("product", ""),
                asset.get("extrainfo", ""),
                asset.get("scan_time", "")
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border

        # 调整列宽
        column_widths = [15, 8, 10, 8, 15, 20, 15, 20, 20]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        # 保存到BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # 返回流式响应
        filename = f"asset_scan_{ip}_{int(time.time())}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
        )

    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出Excel失败：{str(e)}")
