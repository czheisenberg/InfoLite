# -*- coding: utf-8 -*-
# @Desc: 子域名扫描API路由
from fastapi import APIRouter, Query, Depends, HTTPException
from fastapi.responses import JSONResponse, StreamingResponse
from io import BytesIO
import time

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from subdomain_scan import subdomain_scan
from db_operate import insert_subdomain_to_mysql, insert_subdomain_to_es, query_subdomain_from_db
from api.deps import get_db

router = APIRouter(prefix="/api/subdomain", tags=["子域名扫描"])


@router.get("/scan", summary="子域名扫描接口")
async def scan_subdomain(
    domain: str = Query(..., description="目标域名，如 example.com"),
    scan_type: str = Query("all", description="扫描类型：crtsh、chaziyu、dict、all"),
    max_workers: int = Query(50, description="字典爆破最大并发数"),
    refresh: bool = Query(False, description="是否强制刷新，True=重新扫描，False=优先查缓存"),
    save_to_db: bool = Query(True, description="是否保存结果到数据库"),
    current_user: dict = Depends(lambda: None),
    db: dict = Depends(get_db)
):
    """
    子域名扫描接口：
    - refresh=False：优先从数据库查询缓存，有有效数据则直接返回
    - refresh=True：强制重新扫描，并更新数据库
    """
    try:
        if not domain or len(domain) < 3:
            raise HTTPException(status_code=400, detail="请输入有效的域名")

        if "." not in domain:
            raise HTTPException(status_code=400, detail="域名格式不正确")

        domain = domain.strip().lower()
        if domain.startswith("www."):
            domain = domain[4:]

        if not refresh:
            cache_result = query_subdomain_from_db(
                root_domain=domain,
                mysql_conn=db["mysql_conn"],
                mysql_cursor=db["mysql_cursor"],
                es=db["es_conn"],
                scan_expire_time=db["scan_expire_time"]
            )
            if cache_result["is_valid"] and cache_result["data"]:
                alive_count = sum(1 for r in cache_result["data"] if r.get("status") == "alive")
                return JSONResponse(
                    status_code=200,
                    content={
                        "code": 200,
                        "msg": "查询成功（缓存）",
                        "from_cache": True,
                        "data": {
                            "target": domain,
                            "scan_type": scan_type,
                            "total": len(cache_result["data"]),
                            "alive_count": alive_count,
                            "cost_time": 0,
                            "results": cache_result["data"]
                        }
                    }
                )

        result = subdomain_scan(
            domain=domain,
            scan_type=scan_type,
            max_workers=max_workers
        )

        now = int(time.time())
        now_str = time.strftime("%Y-%m-%d %H:%M:%S")
        for item in result["results"]:
            item["scan_time"] = now
            item["scan_time_str"] = now_str

        if save_to_db and result["results"]:
            insert_subdomain_to_mysql(db["mysql_conn"], db["mysql_cursor"], domain, result["results"])
            insert_subdomain_to_es(db["es_conn"], domain, result["results"])

        return JSONResponse(
            status_code=200,
            content={
                "code": 200,
                "msg": "扫描完成",
                "from_cache": False,
                "data": result
            }
        )

    except HTTPException as e:
        return JSONResponse(
            status_code=e.status_code,
            content={"code": e.status_code, "msg": e.detail, "data": None}
        )
    except Exception as e:
        return JSONResponse(
            status_code=500,
            content={"code": 500, "msg": f"服务器内部错误：{str(e)}", "data": None}
        )


@router.get("/export-excel", summary="导出子域名Excel接口")
async def export_subdomain_excel(
    domain: str = Query(..., description="目标域名"),
    current_user: dict = Depends(lambda: None),
    db: dict = Depends(get_db)
):
    """
    导出指定域名的子域名扫描结果为Excel文件
    """
    try:
        # 查询数据库中的子域名数据
        cache_result = query_subdomain_from_db(
            root_domain=domain,
            mysql_conn=db["mysql_conn"],
            mysql_cursor=db["mysql_cursor"],
            es=db["es_conn"],
            scan_expire_time=db["scan_expire_time"]
        )
        results = cache_result["data"] if cache_result["is_valid"] else []

        if not results:
            raise HTTPException(status_code=404, detail=f"域名 {domain} 没有可导出的扫描数据")

        # 创建工作簿
        wb = Workbook()
        ws = wb.active
        ws.title = "子域名扫描结果"

        # 设置表头样式
        header_font = Font(name='微软雅黑', bold=True, size=11, color="FFFFFF")
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_alignment = Alignment(horizontal="center", vertical="center")

        # 边框样式
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # 写入表头
        headers = ["根域名", "子域名", "IP地址", "来源", "状态"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment
            cell.border = thin_border

        # 写入数据
        data_font = Font(name='微软雅黑', size=10)
        data_alignment = Alignment(horizontal="center", vertical="center")

        for row, item in enumerate(results, 2):
            row_data = [
                domain,
                item.get("domain", ""),
                ", ".join(item.get("ips", [])) if isinstance(item.get("ips"), list) else item.get("ips", ""),
                item.get("source", ""),
                item.get("status", "")
            ]
            for col, value in enumerate(row_data, 1):
                cell = ws.cell(row=row, column=col, value=value)
                cell.font = data_font
                cell.alignment = data_alignment
                cell.border = thin_border

        # 调整列宽
        column_widths = [20, 30, 30, 20, 10]
        for col, width in enumerate(column_widths, 1):
            ws.column_dimensions[chr(64 + col)].width = width

        # 保存到BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        # 返回流式响应
        filename = f"subdomain_{domain}_{int(time.time())}.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename*=UTF-8''{filename}"}
        )

    except HTTPException as e:
        raise e
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"导出Excel失败：{str(e)}")
